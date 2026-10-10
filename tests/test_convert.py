"""app.convert.converter - Word to PDF with LibreOffice. Runs only where
soffice is installed (DocxConvertFunction's image, or SOFFICE_PATH locally)."""
from __future__ import annotations

import io

import pytest
from pypdf import PdfReader

from app.convert.converter import ConversionError, convert_to_pdf, find_soffice

docx = pytest.importorskip("docx")
needs_soffice = pytest.mark.skipif(find_soffice() is None, reason="LibreOffice (soffice) isn't installed")


def _word_file() -> bytes:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    document = docx.Document()
    for text, align in [("Centred heading", WD_ALIGN_PARAGRAPH.CENTER), ("Right aligned date", WD_ALIGN_PARAGRAPH.RIGHT),
                        ("Justified body text " * 30, WD_ALIGN_PARAGRAPH.JUSTIFY)]:
        document.add_paragraph(text).alignment = align
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(1, 1).text = "Cell A", "Cell D"
    document.add_page_break()
    document.add_paragraph("Second page")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


@needs_soffice
def test_a_word_file_keeps_its_pages_and_text():
    pages = PdfReader(io.BytesIO(convert_to_pdf(_word_file(), "letter.docx"))).pages
    assert len(pages) == 2
    first = pages[0].extract_text()
    assert "Centred heading" in first and "Right aligned date" in first and "Cell D" in first
    assert "Second page" in pages[1].extract_text()


def test_a_file_that_isnt_really_word_says_so():
    # LibreOffice would import it as plain text and "convert" the garbage.
    with pytest.raises(ConversionError, match="'renamed.docx' isn't a valid Word file"):
        convert_to_pdf(b"not really a docx", "renamed.docx", soffice="unused")


@needs_soffice
def test_a_corrupt_word_file_says_so():
    with pytest.raises(ConversionError, match="broken.docx"):
        convert_to_pdf(_word_file()[:2000], "broken.docx")


def test_only_word_files_are_accepted():
    with pytest.raises(ConversionError, match="isn't a Word document"):
        convert_to_pdf(b"%PDF-1.4", "a.pdf", soffice="unused")
