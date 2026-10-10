"""The technical presentation (app.reports.presentation): which rows get a
deck, what its fixed slides say (verified data only), that the PDF and
PowerPoint renderings match slide for slide, the AI step that drafts the
slides in between (a fake provider - no network), and how a bid pack build
carries the deck into the pack and offers the .pptx for download."""
from __future__ import annotations

import io
import json
import types

import mongomock
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pypdf import PdfReader

from app.api import main as api
from app.intelligence.bid_drafter import BidDraftingError, DraftedPresentation, draft_presentation
from app.reports.bid_generator import CHECKLIST_VERSION, generate_bid_package, submission_item
from app.reports.presentation import build_slides, is_presentation_row, render_presentation

TENDER = {
    "title": "Hiring of Social Media Agency - Content Creation, Marathi, Facebook, Instagram, YouTube",
    "organisation": "Social Justice Department",
    "tender_ref": "57768929",
    "document_summary": {"key_dates": ["Tender Opening: 15/10/2026", "Presentation: 16/10/2026"],
                         "summary_text": "Social media campaign on Facebook and YouTube."},
}
HEADER = {"bid_number": "GEM/2026/B/8097271", "organisation": "Social Justice Department"}

PROFILE = {
    "legal_name": "TEST COMPANY PRIVATE LIMITED",
    "trade_name": "Test Co",
    "former_name": "Old Name LLP",
    "cin": "U00000TS2025PTC000001",
    "registered_office": "1 Test Street, Hyderabad",
    "email": "info@test.example", "phone": "+91-9000000000",
    "authorized_signatory": {"name": "Jane Doe", "designation": "Director"},
    "registrations": [{"name": "UDYAM (MSME) Registration", "number": "UDYAM-TS-00-0000001"},
                      {"name": "DPIIT Start-up Recognition", "certificate_no": "DIPP000001"}],
    "certifications": [{"name": "ISO 9001:2015 (Quality)", "certificate_no": "Q-1"},
                       {"name": "CMMI Maturity Level 3", "certificate_no": "CM-9"}],
    "to_be_verified": ["CMMI certificate CM-9 prints a date after its own expiry"],
    "financials": {
        "annual_turnover": [{"financial_year": "2024-25", "amount_inr": 52389823},
                            {"financial_year": "2025-26", "amount_inr": 35946518}],
        "average_turnover": {"period": "FY 2024-25 to FY 2025-26", "amount_inr": 44168170.5},
        "ca_certificate": {"firm": "Test & Co", "udin": "UDIN123"},
        "balance_sheet": [{"financial_year": "2025-26", "as_at": "2026-03-31", "entity": "TEST COMPANY",
                           "net_worth_inr": 117733410},
                          {"financial_year": "2024-25", "as_at": "2025-03-31", "entity": "Old Name LLP (predecessor)",
                           "net_worth_inr": 4465830}],
    },
    "past_experience": [
        {"client": "Dept of Agriculture, Govt of MP", "short_name": "MP CAEC", "scope_label": "Social Media Campaign",
         "work_order_no": "WO/011", "work_order_date": "2026-02-23", "value_inr": 30250000,
         "value_display": "Rs. 3,02,50,000", "status": "Ongoing", "tags": ["social media", "facebook"],
         "library_document": "Work Order - MP"},
        {"client": "Vidura", "short_name": "Vidura", "value_inr_lakh": 15.0, "status": "Ongoing", "tags": []},
    ],
}

DRAFT = {
    "deck_title": "Social Media Promotion of Schemes",
    "subtitle": "Content · Campaigns · Analytics",
    "focus": ["Facebook", "YouTube", "Marathi"],
    "closing_line": "We look forward to working with the Department",
    "presentation_date": "",
    "slides": [
        {"layout": "cards", "title": "Our understanding", "items": [
            {"heading": "Content", "text": "Marathi creatives for each scheme"},
            {"heading": "Campaigns", "text": "Posting and paid boosting"},
            {"heading": "Reporting", "text": "Weekly reports"}], "notes": "Scope in one view."},
        {"layout": "timeline", "title": "Work plan", "unit": "Day", "span": 30, "phases": [
            {"label": "Kick-off", "start": 1, "end": 1}, {"label": "Campaign execution", "start": 4, "end": 30}],
         "milestones": [{"at": "Day 30", "text": "Final report"}]},
        {"layout": "process", "title": "Approval flow", "items": [
            {"heading": "Draft", "text": "Script"}, {"heading": "Review", "text": "QA"}, {"heading": "Publish", "text": "Live"}]},
        {"layout": "columns", "title": "Platforms", "items": [
            {"heading": "Facebook", "bullets": ["Explainers"]}, {"heading": "YouTube", "bullets": ["Videos"]}]},
        {"layout": "mystery", "title": "Team", "items": [{"heading": "Project manager", "text": "Single contact"}]},
        {"layout": "cards", "title": "   "},
    ],
    "open_items": ["Confirm the presentation venue"],
}


def test_is_presentation_row():
    assert is_presentation_row({"source": "draft", "document": "Technical Presentation"})
    assert is_presentation_row({"source": "draft", "document": "PPT on Approach & Methodology"})
    assert not is_presentation_row({"source": "upload", "document": "Technical Presentation"})
    assert not is_presentation_row({"source": "draft", "document": "Covering Letter"})


def test_fixed_slides_come_from_verified_data_only():
    slides = build_slides(TENDER, PROFILE, HEADER, DraftedPresentation.model_validate(DRAFT))
    kinds = [s.kind for s in slides]
    assert kinds[:3] == ["title", "company", "experience"] and kinds[-1] == "closing"
    # Drafted slides in between - blank titles dropped, unknown layouts drawn as cards.
    assert [s.title for s in slides[3:-1]] == ["Our understanding", "Work plan", "Approval flow", "Platforms", "Team"]
    assert slides[7].kind == "cards"

    title = slides[0]
    assert title.title == "Social Media Promotion of Schemes"
    assert "Tender No. 57768929  |  Bid No. GEM/2026/B/8097271" in title.data["lines"]
    assert title.data["lines"][-1].endswith("16-10-2026")  # the presentation date from the key dates

    company = slides[1]
    assert ("Rs. 4.42 Cr", "Average annual turnover, FY 2024-25 to FY 2025-26 (CA certified)") in company.data["stats"]
    assert ("Rs. 11.77 Cr", "Net worth as at 31-03-2026") in company.data["stats"]
    assert ("Rs. 3.03 Cr", "Combined value of 1 work orders on record") in company.data["stats"]
    chips = dict(company.data["chips"])
    assert chips["ISO 9001:2015"] == "Certificate Q-1"
    assert not any("CMMI" in name for name in chips)  # questioned in to_be_verified - left off
    assert "FY 2024-25 figures are of Old Name LLP, the predecessor entity." in company.data["footnote"]

    rows = slides[2].data["rows"]
    assert len(rows) == 1  # a project with no work order in the library isn't listed
    assert rows[0] == ["Dept of Agriculture, Govt of MP", "WO/011\n23-02-2026", "Social Media Campaign",
                       "3,02,50,000", "Ongoing"]


def test_deck_without_drafted_slides_has_a_placeholder():
    slides = build_slides(TENDER, PROFILE, HEADER, None)
    assert [s.kind for s in slides] == ["title", "company", "experience", "placeholder", "closing"]
    assert slides[0].title == TENDER["title"]
    assert "[TO BE VERIFIED" in slides[3].data["text"]


def test_pdf_and_pptx_match_slide_for_slide(tmp_path):
    from PIL import Image

    logo = tmp_path / "logo.png"
    Image.new("RGB", (260, 100), "navy").save(logo)
    deck = render_presentation(TENDER, PROFILE, HEADER, DraftedPresentation.model_validate(DRAFT), logo)
    assert deck.slide_count == 9

    pdf = PdfReader(io.BytesIO(deck.pdf))
    assert len(pdf.pages) == 9
    assert float(pdf.pages[0].mediabox.width) == pytest.approx(13.333 * 72)
    text = "\n".join(p.extract_text() for p in pdf.pages)
    for phrase in ("Our understanding", "Kick-off", "Rs. 3,02,50,000", "Jane Doe, Director", "Day 30"):
        assert phrase in text

    prs = Presentation(io.BytesIO(deck.pptx))
    assert len(prs.slides) == 9
    texts = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
    assert "Our understanding" in texts[3] and "Marathi creatives for each scheme" in texts[3]
    assert prs.slides[3].notes_slide.notes_text_frame.text == "Scope in one view."


def test_render_without_a_logo(tmp_path):
    deck = render_presentation(TENDER, {"legal_name": "X"}, None, None, tmp_path / "missing.png")
    assert len(PdfReader(io.BytesIO(deck.pdf)).pages) == deck.slide_count == 3  # title, placeholder, closing


class FakeProvider:
    def __init__(self, response):
        self.response = response
        self.calls = []

    def generate(self, system_prompt, user_prompt):
        self.calls.append((system_prompt, user_prompt))
        return self.response


def test_draft_presentation_parses_the_slides():
    provider = FakeProvider(json.dumps(DRAFT))
    row = {"id": "r1", "document": "Technical Presentation", "what_to_upload": "Work plan and social media strategy"}
    deck = draft_presentation(TENDER, row, PROFILE, [], provider)
    assert deck.deck_title == "Social Media Promotion of Schemes"
    assert len(deck.slides) == 5  # the blank-titled slide is dropped
    assert deck.slides[4].layout == "mystery"
    system, user = provider.calls[0]
    assert "NEVER INVENT A FACT" in system
    assert "Presentation asked for: \"Technical Presentation\" - Work plan and social media strategy" in user
    assert "Presentation: 16/10/2026" in user
    assert "MP CAEC" in user


@pytest.mark.parametrize("response", ["not json", json.dumps({"slides": []}), json.dumps({"slides": "x"})])
def test_draft_presentation_failures_raise(response):
    with pytest.raises(BidDraftingError):
        draft_presentation(TENDER, {"document": "Technical Presentation"}, PROFILE, [], FakeProvider(response))


def _checklist(*rows):
    return {"version": CHECKLIST_VERSION, "header": HEADER, "items": list(rows), "notes": []}


def test_bid_pack_carries_the_slides_at_the_presentation_row():
    letter = submission_item("Covering Letter", source="draft", status="To be prepared")
    deck_row = submission_item("Technical Presentation", "Work plan", "Technical Presentation", source="draft",
                               status="To be prepared")
    deck = render_presentation(TENDER, PROFILE, HEADER, DraftedPresentation.model_validate(DRAFT))
    pack = PdfReader(io.BytesIO(generate_bid_package(
        TENDER, [], PROFILE, [], checklist=_checklist(letter, deck_row), cover_page=False,
        presentations={deck_row["id"]: deck.pdf})))
    widths = [round(float(p.mediabox.width)) for p in pack.pages]
    assert widths.count(960) == deck.slide_count  # the slides, at their own 16:9 size
    # ...right after the checklist page and the covering letter's page(s).
    first_slide = widths.index(960)
    assert "TECHNICAL PRESENTATION" in pack.pages[first_slide].extract_text()


@pytest.fixture()
def built(monkeypatch):
    """_build_bid_pack run for real on a mongomock tender with one
    presentation row - AI calls and file storage stubbed."""
    tenders = mongomock.MongoClient().db["tenders"]
    row = submission_item("Technical Presentation", "Work plan", "Technical Presentation", source="draft",
                          status="To be prepared")
    tender_id = tenders.insert_one({**TENDER, "checklist": _checklist(row)}).inserted_id
    files: dict[str, bytes] = {}
    monkeypatch.setattr(api, "get_collection", lambda: tenders)
    monkeypatch.setattr(api.storage, "load_document_text", lambda t: t)
    monkeypatch.setattr(api.storage, "put_bytes", lambda key, data, content_type=None: files.__setitem__(key, data))
    monkeypatch.setattr(api.storage, "get_bytes", lambda key: files.get(key))
    monkeypatch.setattr(api.storage, "delete", lambda key: files.pop(key, None))
    monkeypatch.setattr(api.storage, "download_url", lambda key, filename: None)
    monkeypatch.setattr(api, "_generation_inputs", lambda: ([], PROFILE, []))
    monkeypatch.setattr(api, "draft_presentation",
                        lambda *a, **kw: DraftedPresentation.model_validate(DRAFT))

    def no_letters(*a, **kw):
        raise AssertionError("a presentation row is not drafted as a letter")

    monkeypatch.setattr(api, "draft_checklist_documents", no_letters)
    object_id, tender = api._find_tender(str(tender_id))
    api._build_bid_pack(object_id, tender)
    return types.SimpleNamespace(tenders=tenders, id=str(tender_id), files=files, row=row,
                                 client=TestClient(api.app))


def test_build_stores_the_pptx_beside_the_pack(built):
    tender = built.tenders.find_one({})
    assert tender["presentation_key"].endswith(".pptx") and tender["bid_pack_key"].endswith(".pdf")
    assert len(Presentation(io.BytesIO(built.files[tender["presentation_key"]])).slides) == 9
    item = tender["checklist"]["items"][0]
    assert item["status"] == "Enclosed"
    notes = [n["requirement"] for n in tender["checklist"]["notes"]]
    assert "Technical Presentation: Confirm the presentation venue" in notes

    serialized = built.client.get(f"/tenders/{built.id}/generate-bid").json()["tender"]
    assert serialized["presentation_available"] is True and "presentation_key" not in serialized
    link = built.client.get(f"/tenders/{built.id}/presentation/link").json()
    assert link == {"url": None, "filename": "presentation-57768929.pptx"}
    res = built.client.get(f"/tenders/{built.id}/presentation")
    assert res.status_code == 200 and res.content == built.files[tender["presentation_key"]]
    assert "presentation-57768929.pptx" in res.headers["content-disposition"]


def test_rebuilding_replaces_the_previous_pptx(built):
    first = built.tenders.find_one({})["presentation_key"]
    object_id, tender = api._find_tender(built.id)
    api._build_bid_pack(object_id, tender)
    second = built.tenders.find_one({})["presentation_key"]
    assert second != first and first not in built.files and second in built.files


def test_no_presentation_link_without_one(built):
    built.tenders.update_one({}, {"$set": {"presentation_key": None}})
    assert built.client.get(f"/tenders/{built.id}/presentation/link").status_code == 404
