"""app.convert.word_text and POST /stamp-document/word/... - "Edit text"
for a Word file on the Sign & Stamp tab: words, formatting, new and deleted
paragraphs - everything else in the file kept."""
from __future__ import annotations

import io

import pytest

from app.convert.word_text import NewParagraph, ParagraphChange, WordTextError, apply_edits, read_paragraphs
from tests.test_stamp_document import _pdf, _upload_in_parts, client, transfers  # noqa: F401 - fixtures

docx = pytest.importorskip("docx")


def _word_file() -> bytes:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    document = docx.Document()
    heading = document.add_paragraph("Covering Letter")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    body = document.add_paragraph()
    body.add_run("Turnover:").bold = True
    body.add_run(" approximately ")
    body.add_run("7 Crore").bold = True
    document.add_paragraph("")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Cell text"
    document.sections[0].footer.paragraphs[0].text = "Footer line"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _ids(content: bytes) -> dict[str, str]:
    return {p["text"]: p["id"] for p in read_paragraphs(content)}


def _body(content: bytes):
    return docx.Document(io.BytesIO(content)).paragraphs


def test_every_paragraph_is_listed_with_its_formatting():
    listed = read_paragraphs(_word_file())
    rows = [(p["where"], p["text"], p["bold"], p["align"]) for p in listed]
    assert rows[:3] == [("Body", "Covering Letter", False, "center"),
                        ("Body", "Turnover: approximately 7 Crore", None, None),  # partly bold
                        ("Body", "", False, None)]  # blank lines too
    assert ("Table", "Cell text", False, None) in rows and ("Footer", "Footer line", False, None) in rows


def test_an_edit_keeps_each_untouched_characters_formatting():
    src = _word_file()
    ids = _ids(src)
    out = apply_edits(src, {ids["Turnover: approximately 7 Crore"]: ParagraphChange(text="Turnover: about 9 Crore"),
                            ids["Covering Letter"]: ParagraphChange(text="Covering Letter for Bid")})
    turnover = next(p for p in _body(out) if p.text.startswith("Turnover"))
    assert [(r.text, bool(r.bold)) for r in turnover.runs] == [("Turnover:", True), (" about ", False), ("9 Crore", True)]
    assert _body(out)[0].text == "Covering Letter for Bid" and _body(out)[0].alignment == 1  # still centred


def test_formatting_alignment_and_size_change_the_whole_paragraph():
    src = _word_file()
    out = apply_edits(src, {_ids(src)["Turnover: approximately 7 Crore"]: ParagraphChange(
        bold=False, italic=True, underline=True, align="justify", size=13)})
    listed = next(p for p in read_paragraphs(out) if p["text"].startswith("Turnover"))
    assert (listed["bold"], listed["italic"], listed["underline"], listed["align"], listed["size"]) == \
        (False, True, True, "justify", 13.0)


def test_a_blank_line_can_be_typed_into():
    src = _word_file()
    blank = next(p["id"] for p in read_paragraphs(src) if p["text"] == "")
    out = apply_edits(src, {blank: ParagraphChange(text="Now with words", bold=True)})
    para = _body(out)[2]
    assert para.text == "Now with words" and all(r.bold for r in para.runs)


def test_paragraphs_can_be_added_below_and_deleted():
    src = _word_file()
    ids = _ids(src)
    out = apply_edits(
        src,
        {ids["Covering Letter"]: ParagraphChange(delete=True)},
        [NewParagraph(after=ids["Turnover: approximately 7 Crore"], key="n1", text="Added line"),
         NewParagraph(after="n1", text="Added after that", align="right")],
    )
    texts = [p.text for p in _body(out)]
    assert texts[:3] == ["Turnover: approximately 7 Crore", "Added line", "Added after that"]
    assert _body(out)[1].runs[0].bold  # the new line takes its anchor's first formatting
    assert _body(out)[2].alignment == 2  # right


def test_the_only_paragraph_of_a_table_cell_is_emptied_not_removed():
    src = _word_file()
    out = apply_edits(src, {_ids(src)["Cell text"]: ParagraphChange(delete=True)})
    cell = docx.Document(io.BytesIO(out)).tables[0].cell(0, 0)
    assert len(cell.paragraphs) == 1 and cell.text == ""


def test_bad_changes_and_non_docx_files_are_refused():
    src = _word_file()
    with pytest.raises(WordTextError):
        apply_edits(src, {"word/document.xml#999": ParagraphChange(text="x")})
    with pytest.raises(WordTextError):
        apply_edits(src, {"../evil.xml#1": ParagraphChange(text="x")})
    with pytest.raises(WordTextError):
        apply_edits(src, inserts=[NewParagraph(after="nowhere", text="x")])
    with pytest.raises(WordTextError, match=".docx"):
        read_paragraphs(b"\xd0\xcf\x11\xe0 old .doc")


def test_edit_text_round_trip_through_the_api(client, transfers, monkeypatch):  # noqa: F811
    monkeypatch.delenv("DOCX_CONVERT_FUNCTION_NAME", raising=False)
    converted = []
    monkeypatch.setattr("app.api.main.convert_to_pdf", lambda content, filename: converted.append(content) or _pdf(1))
    parts = _upload_in_parts(client, _word_file())
    upload = {"upload_id": "a" * 32, "parts": parts, "filename": "Letter.docx"}
    assert client.post("/stamp-document/convert", json=upload).json()["status"] == "done"

    listed = client.post("/stamp-document/word/paragraphs", json=upload).json()["paragraphs"]
    target = next(p for p in listed if p["text"].startswith("Turnover"))
    job = client.post("/stamp-document/word/edit", json={
        **upload,
        "edits": {target["id"]: {"text": "Turnover: 9 Crore", "align": "center"}},
        "inserts": [{"after": target["id"], "text": "A new line"}],
    }).json()
    assert job["status"] == "done" and job["filename"] == "Letter.pdf"
    texts = [p["text"] for p in read_paragraphs(converted[-1])]
    assert texts[1:3] == ["Turnover: 9 Crore", "A new line"]

    assert client.post("/stamp-document/word/edit", json={**upload, "filename": "Letter.doc"}).status_code == 422
    bad = client.post("/stamp-document/word/edit", json={**upload, "edits": {target["id"]: {"align": "diagonal"}}})
    assert bad.status_code == 422


def test_line_breaks_and_tabs_are_shown_and_kept():
    document = docx.Document()
    para = document.add_paragraph()
    run = para.add_run("Managing Director")
    run.add_break()
    run.add_text("Old Address")
    para.add_run("\tTab after")
    buf = io.BytesIO()
    document.save(buf)
    src = buf.getvalue()
    listed = read_paragraphs(src)[0]
    assert listed["text"] == "Managing Director\nOld Address\tTab after"

    out = apply_edits(src, {listed["id"]: ParagraphChange(text="Managing Director\nNew Address\nLine 3\tTab after")})
    assert read_paragraphs(out)[0]["text"] == "Managing Director\nNew Address\nLine 3\tTab after"
    joined = apply_edits(src, {listed["id"]: ParagraphChange(text="Managing Director, Old Address\tTab after")})
    assert read_paragraphs(joined)[0]["text"] == "Managing Director, Old Address\tTab after"
